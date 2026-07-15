"""
Minimal HTML to Markdown converter using trafilatura.
Scrapes policy research documents and saves as clean Markdown files.

Install: pip install trafilatura

Usage: python 00_html_to_markdown.py
"""

import json
import re
from pathlib import Path
from urllib.parse import urlparse

import openpyxl
from trafilatura import fetch_url, extract

from config import PROJECT_ROOT, METADATA_PATH
from utils.utils import derive_subfolder, slugify

URLS_CONFIG_PATH = PROJECT_ROOT/"urls_config.json"
OUTPUT_BASE_DIR = PROJECT_ROOT/"data"/"md"

# Delimiter for splitting Upward Mobility Framework predictor pages.
# These 24 pages are foundational, synthesized docs for this KB. Each contains
# an "Evidence" section and a "Promising Local Policy Interventions" section
# that map to separate document_type metadata values in the knowledge base.
FRAMEWORK_SPLIT_DELIMITER = "#### Promising Local Policy Interventions"
def split_framework_doc(content, base_filename, evidence_doc_type, interventions_doc_type):
    """Split a Framework predictor page into separate evidence and interventions docs.

    Only applies to the 24 predictor pages. Each page has a consistent structure:
    frontmatter + intro + Evidence section, then a Policy Interventions section.
    We split on FRAMEWORK_SPLIT_DELIMITER so each half becomes its own document.

    Args:
        content: Full markdown content (with YAML frontmatter) from trafilatura
        base_filename: Original filename (e.g., "jobs_paying_living_wages.md")
        evidence_doc_type: document_type value for the evidence half
            (must match the Information Type in the metadata spreadsheet)
        interventions_doc_type: document_type value for the interventions half
            (must match the Information Type in the metadata spreadsheet)

    Returns:
        List of (content, filename, document_type) tuples — one for evidence,
        one for interventions.
    """
    # Extract frontmatter block (everything between the --- delimiters)
    _, yaml_fields, body = content.split('---', 2)
    frontmatter = f"---{yaml_fields}---"

    # Split the body on the delimiter
    evidence_body, interventions_body = body.split(FRAMEWORK_SPLIT_DELIMITER)

    evidence_content = f"{frontmatter}{evidence_body}".rstrip()
    interventions_content = f"{frontmatter}\n{FRAMEWORK_SPLIT_DELIMITER}{interventions_body}".rstrip()

    evidence_filename = base_filename.replace('.md', '_evidence.md')
    interventions_filename = base_filename.replace('.md', '_interventions.md')

    return [
        (evidence_content, evidence_filename, evidence_doc_type),
        (interventions_content, interventions_filename, interventions_doc_type),
    ]


def generate_urls_config():
    """Read metadata spreadsheet and generate urls_config.json for HTML scraping.

    Filters to scrapable HTML rows (non-paywalled) and writes a config file
    with URL and subfolder for each document.
    """
    wb = openpyxl.load_workbook(METADATA_PATH)
    ws = wb["Metadata"]
    headers = [cell.value for cell in ws[1]]

    urls = []
    skipped = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        row_data = dict(zip(headers, row))
        if row_data['File Type'] != 'HTML' or row_data['Paywall?'] == 'Yes':
            skipped += 1
            continue

        subfolder = derive_subfolder(row_data['Information Source'])

        urls.append(
            {'url': row_data['Document URL'],
             'subfolder': subfolder}
        )

    config = {"urls": urls}
    URLS_CONFIG_PATH.write_text(json.dumps(config, indent=2))
    print(f"Wrote {len(urls)} URLs to {URLS_CONFIG_PATH} ({skipped} rows skipped)")


def load_config(config_path=URLS_CONFIG_PATH):
    """Load URL configuration from JSON file."""
    with open(config_path) as f:
        return json.load(f)


def scrape_url(url):
    """
    Fetch URL and extract content as markdown.

    Args:
        url: Web page URL to scrape

    Returns:
        Extracted markdown content or None on failure
    """
    print(f"  Fetching: {url}")
    downloaded = fetch_url(url)

    if not downloaded:
        print(f"  Failed to download: {url}")
        return None

    # Extract with markdown format and metadata
    # favor_precision=True for cleaner output, include_tables for data preservation
    result = extract(
        downloaded,
        output_format="markdown",
        with_metadata=True,
        include_tables=True,
        include_formatting=True,
        favor_precision=True,
    )

    if not result:
        print(f"  No content extracted from: {url}")
        return None

    return result


def extract_title(content):
    """Extract the title from trafilatura's YAML frontmatter.

    Args:
        content: Markdown string with YAML frontmatter delimited by '---'

    Returns:
        Title string, or None if not found
    """
    # Search only within the title YAML using the '---' dividers
    content = content.split('---')[1]
    match = re.search('title: ', content)
    if not match:
        return None
    start_pos = match.end()
    end_pos = start_pos + content[start_pos:].find('\n')
    return content[start_pos:end_pos]


def derive_filename(content, url):
    """Derive a .md filename from scraped content, falling back to URL path.

    Args:
        content: Scraped markdown content (with YAML frontmatter)
        url: Original URL (used as fallback)

    Returns:
        Filename string ending in .md
    """
    title = extract_title(content)
    if title:
        return f"{slugify(title)}.md"

    # Fallback: use the last segment of the URL path
    path = urlparse(url).path.strip('/').split('/')[-1]
    return f"{slugify(path)}.md"


def save_markdown(content, filename, subfolder=None):
    """Save content to markdown file in output directory with optional subfolder."""
    if subfolder:
        output_dir = OUTPUT_BASE_DIR / subfolder
    else:
        output_dir = OUTPUT_BASE_DIR
    output_dir.mkdir(parents=True, exist_ok=True)
    filepath = output_dir / filename
    filepath.write_text(content, encoding="utf-8")
    print(f"  Saved: {filepath}")


def save_config(config, config_path=URLS_CONFIG_PATH):
    """Write updated config back to JSON file."""
    with open(config_path, 'w') as f:
        json.dump(config, f, indent=2)


def main():
    """Process all URLs from config and save as markdown files."""
    generate_urls_config()
    config = load_config()

    print(f"Processing {len(config['urls'])} URLs...\n")

    final_items = []

    for item in config["urls"]:
        url = item["url"]
        subfolder = item.get("subfolder")

        print(f"[{url}]")

        content = scrape_url(url)
        if not content:
            print(f"  Skipped: no content")
            print()
            final_items.append(item)  # keep so utils.py can flag as scrape failure
            continue

        base_filename = derive_filename(content, url)

        if FRAMEWORK_SPLIT_DELIMITER in content:
            splits = split_framework_doc(
                content, base_filename,
                evidence_doc_type="Predictor Page Evidence",
                interventions_doc_type="Predictor Page Solutions",
            )
            # Splits replace the original entry — don't append `item`.
            # Its {url, subfolder} are already carried by each split entry,
            # so keeping the original would show up as a phantom scrape failure.
            for split_content, split_filename, doc_type in splits:
                print(f"  Split: {split_filename} ({doc_type})")
                save_markdown(split_content, split_filename, subfolder)
                final_items.append({
                    "url": url,
                    "subfolder": subfolder,
                    "filename": split_filename,
                    "document_type": doc_type,
                })
        else:
            item["filename"] = base_filename
            print(f"  Filename: {base_filename}")
            save_markdown(content, base_filename, subfolder)
            final_items.append(item)

        print()

    config["urls"] = final_items

    # Write filenames back to config so downstream scripts can use them
    save_config(config)
    print("Done!")


if __name__ == "__main__":
    main()
